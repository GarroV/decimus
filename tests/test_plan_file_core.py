"""Файл «туда и обратно» электронного плана — ядро без базы (#1168, шаг 3).

Что здесь держится: какие строки файла ложатся в план, а какие нет и почему (чужой id,
повтор, слишком длинный текст), какой срок принимается (не позже положенного, не раньше
сегодня, дата), что файл — граница системы (формулы не исполняются, мусор и бомбы не
разбираются, чужой шаблон не принимается). Запись в базу и конфликт ревизий проверяет
`tests/test_db_plan_file.py`.
"""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, date, datetime, timedelta
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from src.db import action_plans as plans
from src.db.plan_file import (
    DUE_NOT_DATE,
    SKIP_DUPLICATE,
    SKIP_TOO_LONG,
    SKIP_UNKNOWN,
    FileRow,
    reconcile,
)
from src.db.plan_rows import MAX_MEASURE, PlanRow
from src.domain.plan_due import OWN_DUE_LATE, OWN_DUE_PAST, own_due_problem
from src.web import plan_sheet
from src.web.plan_sheet import SheetError

ЗАПРОС = "11111111-1111-1111-1111-111111111111"
СТРОКА_1 = "55555555-5555-5555-5555-555555555551"
СТРОКА_2 = "55555555-5555-5555-5555-555555555552"
НАЧАЛО = date(2026, 10, 4)
СЕГОДНЯ = date(2026, 10, 6)


def строка(ident: str = СТРОКА_1, position: int = 1, **kw: Any) -> PlanRow:
    данные: dict[str, Any] = {
        "id": ident,
        "request_id": ЗАПРОС,
        "finding_id": "66666666-6666-6666-6666-666666666666",
        "position": position,
        "item_code": "PRD01",
        "level": "D2",
        "zone": "fridge",
        "title_ru": "Заготовка размечена",
        "title_en": "Prep item is tagged",
        "days": 5,
        "measure": "",
        "responsible": "",
        "reason": "",
        "own_due": None,
        "extended_due": None,
        "due_on": НАЧАЛО + timedelta(days=5),
        "status": "filling",
        "return_reason": "",
        "updated_at": datetime(2026, 10, 4, tzinfo=UTC),
    }
    данные.update(kw)
    return PlanRow(**данные)


def из_файла(row_id: str = СТРОКА_1, line: int = 11, **kw: Any) -> FileRow:
    данные: dict[str, Any] = {
        "line": line,
        "row_id": row_id,
        "measure": "Переразметим",
        "responsible": "Шеф",
        "reason": "",
        "due": None,
    }
    данные.update(kw)
    return FileRow(**данные)


def сверить(file_rows: list[FileRow], rows: list[PlanRow] | None = None) -> Any:
    return reconcile(file_rows, rows or [строка()], count_from=НАЧАЛО, today=СЕГОДНЯ)


# --- правило своего срока ----------------------------------------------------


def test_срок_позже_положенного_не_принимается_и_называет_границу() -> None:
    проблема = own_due_problem(
        НАЧАЛО + timedelta(days=6),
        None,
        count_from=НАЧАЛО,
        days=5,
        extended=None,
        today=СЕГОДНЯ,
    )

    assert проблема == (OWN_DUE_LATE, НАЧАЛО + timedelta(days=5))


def test_срок_раньше_сегодня_не_принимается() -> None:
    проблема = own_due_problem(
        НАЧАЛО, None, count_from=НАЧАЛО, days=5, extended=None, today=СЕГОДНЯ
    )

    assert проблема == (OWN_DUE_PAST, СЕГОДНЯ)


def test_перенос_с_согласием_двигает_предел_а_прежний_срок_не_проверяется_заново() -> None:
    перенос = НАЧАЛО + timedelta(days=9)
    assert (
        own_due_problem(перенос, None, count_from=НАЧАЛО, days=5, extended=перенос, today=СЕГОДНЯ)
        is None
    )
    прежний = НАЧАЛО  # уже в прошлом, но его не меняют
    assert (
        own_due_problem(прежний, прежний, count_from=НАЧАЛО, days=5, extended=None, today=СЕГОДНЯ)
        is None
    )


# --- сверка файла с планом ---------------------------------------------------


def test_верная_строка_ложится_в_план() -> None:
    итог = сверить([из_файла(reason="Не успели", due=СЕГОДНЯ + timedelta(days=1))])

    assert итог.read == 1 and итог.skipped == () and итог.due_refused == ()
    (изменение,) = итог.changes
    assert (изменение.measure, изменение.responsible, изменение.reason) == (
        "Переразметим",
        "Шеф",
        "Не успели",
    )
    assert изменение.own_due == СЕГОДНЯ + timedelta(days=1)
    assert изменение.row.id == СТРОКА_1


def test_чужой_id_строки_пропускается_с_причиной_и_ничего_не_меняет() -> None:
    чужой = "99999999-9999-9999-9999-999999999999"

    итог = сверить([из_файла(чужой, line=12)])

    assert итог.changes == ()
    assert [(s.line, s.code) for s in итог.skipped] == [(12, SKIP_UNKNOWN)]


def test_повтор_id_во_второй_раз_пропускается_первый_берётся() -> None:
    итог = сверить([из_файла(measure="Первая"), из_файла(measure="Вторая", line=12)])

    assert [c.measure for c in итог.changes] == ["Первая"]
    assert [(s.line, s.code) for s in итог.skipped] == [(12, SKIP_DUPLICATE)]


def test_слишком_длинный_текст_пропускает_строку_а_не_обрезает() -> None:
    итог = сверить([из_файла(measure="я" * (MAX_MEASURE + 1))])

    assert итог.changes == ()
    assert [s.code for s in итог.skipped] == [SKIP_TOO_LONG]


def test_срок_позже_положенного_отвергается_а_остальные_поля_строки_ложатся() -> None:
    итог = сверить([из_файла(due=НАЧАЛО + timedelta(days=6))])

    (отказ,) = итог.due_refused
    assert (отказ.position, отказ.code, отказ.bound) == (
        1,
        OWN_DUE_LATE,
        НАЧАЛО + timedelta(days=5),
    )
    (изменение,) = итог.changes
    assert изменение.measure == "Переразметим" and изменение.own_due is None


def test_срок_не_дата_отвергается_и_прежний_срок_остаётся() -> None:
    прежний = СЕГОДНЯ + timedelta(days=1)
    итог = сверить(
        [из_файла(due=None, due_unreadable=True)], [строка(own_due=прежний, due_on=прежний)]
    )

    assert [d.code for d in итог.due_refused] == [DUE_NOT_DATE]
    (изменение,) = итог.changes
    assert изменение.own_due == прежний


def test_пустой_срок_в_файле_снимает_свой_срок_как_пустое_поле_формы() -> None:
    прежний = СЕГОДНЯ + timedelta(days=1)

    итог = сверить([из_файла()], [строка(own_due=прежний, due_on=прежний)])

    assert итог.due_refused == ()
    assert итог.changes[0].own_due is None


def test_нет_меры_или_ответственного_называет_номер_строки_плана() -> None:
    итог = сверить(
        [из_файла(СТРОКА_1, measure=""), из_файла(СТРОКА_2, line=12, responsible="  ")],
        [строка(СТРОКА_1, 1), строка(СТРОКА_2, 2)],
    )

    assert итог.unfilled == (1, 2)


def test_строка_плана_которой_нет_в_файле_остаётся_как_на_экране() -> None:
    итог = сверить(
        [из_файла(СТРОКА_1)],
        [строка(СТРОКА_1, 1), строка(СТРОКА_2, 2, measure="Уже заполнено")],
    )

    assert итог.absent == (2,)
    assert [c.row.id for c in итог.changes] == [СТРОКА_1]


def test_совпавшая_с_экраном_строка_не_пишется() -> None:
    уже = строка(measure="Переразметим", responsible="Шеф")

    итог = сверить([из_файла()], [уже])

    assert итог.changes == () and итог.unchanged == 1


# --- выгрузка и разбор -------------------------------------------------------


def запрос_плана(rows: tuple[PlanRow, ...], revision: int = 7) -> plans.PlanRequest:
    return plans.PlanRequest(
        id=ЗАПРОС,
        inspection_id="33333333-3333-3333-3333-333333333333",
        country="GE",
        due_on=date(2026, 10, 14),
        status="requested",
        origin="auto",
        requested_by="garva",
        requested_at=datetime(2026, 10, 4, tzinfo=UTC),
        unit_id="44444444-4444-4444-4444-444444444444",
        unit_name="Batumi-1",
        inspection_date=date(2026, 10, 3),
        pct=91.5,
        grade="C",
        files=(),
        events=(),
        count_from=НАЧАЛО,
        revision=revision,
        electronic=True,
        review_comment=None,
        rows=rows,
    )


def выгрузка(lang: str = "ru", **kw: Any) -> bytes:
    rows = (строка(СТРОКА_1, 1, **kw), строка(СТРОКА_2, 2, level="D1", title_ru="Пол чистый"))
    return plan_sheet.build(запрос_плана(rows), {"fridge": "Холодильник"}, lang)


def правка(data: bytes, ячейки: dict[str, Any]) -> bytes:
    """Файл после того, как партнёр (или злоумышленник) поправил ячейки первого листа."""
    книга = load_workbook(io.BytesIO(data))
    лист = книга.worksheets[0]
    for адрес, значение in ячейки.items():
        if isinstance(значение, str) and any(ord(c) < 32 and c not in "\n\t\r" for c in значение):
            # Библиотека такие знаки не пишет; в чужом файле они могут быть — кладём в обход.
            лист[адрес] = "x"
            лист[адрес]._value = значение
        else:
            лист[адрес] = значение
    out = io.BytesIO()
    книга.save(out)
    return out.getvalue()


def первая_строка_данных(data: bytes) -> int:
    лист = load_workbook(io.BytesIO(data)).worksheets[0]
    return (
        next(r for r in range(1, 40) if лист.cell(row=r, column=10).value == plan_sheet.ID_HEADER)
        + 1
    )


def test_выгрузка_читается_обратно_с_id_строк_и_ревизией() -> None:
    данные = выгрузка()

    разобрано = plan_sheet.parse(данные)

    assert (разобрано.request_id, разобрано.revision) == (ЗАПРОС, 7)
    assert [r.row_id for r in разобрано.rows] == [СТРОКА_1, СТРОКА_2]
    assert all(r.measure == "" and r.due is None for r in разобрано.rows)


def test_что_партнёр_вписал_в_жёлтые_колонки_читается_значениями() -> None:
    данные = выгрузка()
    первая = первая_строка_данных(данные)
    правленый = правка(
        данные,
        {
            f"F{первая}": "Переразметим",
            f"G{первая}": "Шеф",
            f"H{первая}": "Не успели",
            f"I{первая}": date(2026, 10, 8),
            f"I{первая + 1}": "09.10.2026",
        },
    )

    строки = plan_sheet.parse(правленый).rows

    assert (строки[0].measure, строки[0].responsible, строки[0].reason) == (
        "Переразметим",
        "Шеф",
        "Не успели",
    )
    assert строки[0].due == date(2026, 10, 8) and строки[1].due == date(2026, 10, 9)


def test_срок_не_дата_помечается_нечитаемым_а_не_пропадает() -> None:
    данные = выгрузка()
    первая = первая_строка_данных(данные)

    строки = plan_sheet.parse(
        правка(данные, {f"I{первая}": "скоро", f"I{первая + 1}": "31.02.2026"})
    ).rows

    assert [(r.due, r.due_unreadable) for r in строки] == [(None, True), (None, True)]


def test_формула_в_ячейке_не_исполняется() -> None:
    данные = выгрузка()
    первая = первая_строка_данных(данные)

    строки = plan_sheet.parse(
        правка(данные, {f"F{первая}": "=1+1", f"G{первая}": "=SUM(1,2)"})
    ).rows

    # Читается сохранённое значение; у формулы, которую никто не считал, его нет.
    assert (строки[0].measure, строки[0].responsible) == ("", "")


def test_текст_со_знаком_равно_остаётся_текстом_и_при_выгрузке_не_становится_формулой() -> None:
    опасный = '=HYPERLINK("http://evil.example","клик")'
    запрос = запрос_плана((строка(measure=опасный),))

    книга = load_workbook(io.BytesIO(plan_sheet.build(запрос, {}, "ru")))
    ячейка = книга.worksheets[0].cell(
        row=первая_строка_данных(plan_sheet.build(запрос, {}, "ru")), column=6
    )

    assert ячейка.data_type == "s" and ячейка.value == опасный


def test_колонки_системы_закрыты_колонки_партнёра_открыты_id_скрыт() -> None:
    данные = выгрузка()
    лист = load_workbook(io.BytesIO(данные)).worksheets[0]
    первая = первая_строка_данных(данные)

    assert лист.protection.sheet is True
    assert all(лист.cell(row=первая, column=c).protection.locked for c in (1, 2, 3, 4, 5, 10))
    assert not any(лист.cell(row=первая, column=c).protection.locked for c in (6, 7, 8, 9))
    assert лист.column_dimensions["J"].hidden is True


def test_заголовки_на_языке_интерфейса() -> None:
    ru = load_workbook(io.BytesIO(выгрузка("ru"))).worksheets[0]
    en = load_workbook(io.BytesIO(выгрузка("en"))).worksheets[0]
    заголовок = первая_строка_данных(выгрузка("ru")) - 1

    assert ru.cell(row=заголовок, column=6).value == "Что сделаем"
    assert en.cell(row=заголовок, column=6).value == "What we will do"


def test_срок_по_методике_в_файле_это_положенный_срок_строки() -> None:
    данные = выгрузка()
    лист = load_workbook(io.BytesIO(данные)).worksheets[0]

    значение = лист.cell(row=первая_строка_данных(данные), column=5).value

    assert значение.date() == НАЧАЛО + timedelta(days=5)


def test_чужая_книга_без_нашей_метки_не_разбирается() -> None:
    книга = Workbook()
    книга.active["A1"] = "что-то своё"
    out = io.BytesIO()
    книга.save(out)

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(out.getvalue())

    assert exc.value.code == plan_sheet.ERR_NOT_OURS


def test_без_метки_шаблона_книга_не_наша_даже_с_id_запроса_и_ревизией() -> None:
    данные = правка(выгрузка(), {"J1": "чужая метка"})

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(данные)

    assert exc.value.code == plan_sheet.ERR_NOT_OURS


def test_метка_есть_а_ревизии_нет_книга_не_наша() -> None:
    данные = правка(выгрузка(), {"J3": "не число"})

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(данные)

    assert exc.value.code == plan_sheet.ERR_NOT_OURS


@pytest.mark.parametrize("мусор", [b"", b"not a zip at all", b"PK\x03\x04" + b"\x00" * 40])
def test_не_xlsx_не_разбирается(мусор: bytes) -> None:
    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(мусор)

    assert exc.value.code == plan_sheet.ERR_NOT_XLSX


def _с_частью(данные: bytes, имя: str, тело: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(данные)) as src, zipfile.ZipFile(out, "w") as dst:
        for part in src.infolist():
            dst.writestr(part, src.read(part.filename))
        dst.writestr(имя, тело)
    return out.getvalue()


def test_книга_с_внешними_ссылками_не_разбирается() -> None:
    данные = _с_частью(выгрузка(), "xl/externalLinks/externalLink1.xml", b"<externalLink/>")

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(данные)

    assert exc.value.code == plan_sheet.ERR_EXTERNAL


def test_архив_с_огромным_распакованным_объёмом_не_открывается() -> None:
    данные = _с_частью(
        выгрузка(), "xl/media/bomb.bin", b"\x00" * (plan_sheet.MAX_UNPACKED_BYTES + 1)
    )

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(данные)

    assert exc.value.code == plan_sheet.ERR_TOO_BIG


def test_строк_больше_предела_файл_не_разбирается() -> None:
    данные = выгрузка()
    первая = первая_строка_данных(данные)
    ячейки = {
        f"J{первая + i}": f"77777777-7777-7777-7777-{i:012d}"
        for i in range(plan_sheet.MAX_DATA_ROWS + 1)
    }

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(правка(данные, ячейки))

    assert exc.value.code == plan_sheet.ERR_TOO_BIG


def test_управляющие_знаки_в_значении_вычищаются() -> None:
    assert plan_sheet._clean_text("Меры\x07 тут\x00") == "Меры тут"
    assert plan_sheet._clean_text(5.0) == "5" and plan_sheet._clean_text(None) == ""


def test_книга_с_управляющим_знаком_в_xml_не_разбирается_а_не_падает() -> None:
    данные = выгрузка()
    первая = первая_строка_данных(данные)

    with pytest.raises(SheetError) as exc:
        plan_sheet.parse(правка(данные, {f"F{первая}": "Меры\x07 тут"}))

    assert exc.value.code == plan_sheet.ERR_NOT_XLSX


def test_выгрузка_не_падает_на_управляющем_знаке_в_тексте_из_базы() -> None:
    запрос = запрос_плана((строка(measure="Меры\x07 тут"),))

    данные = plan_sheet.build(запрос, {}, "ru")

    assert plan_sheet.parse(данные).rows[0].measure == "Меры тут"


def test_мусор_в_колонке_id_доходит_до_сверки_и_там_без_пары() -> None:
    данные = выгрузка()
    первая = первая_строка_данных(данные)

    строки = plan_sheet.parse(правка(данные, {f"J{первая}": "drop table"})).rows

    assert строки[0].row_id == "drop table"
    итог = reconcile(
        строки, [строка(СТРОКА_1, 1), строка(СТРОКА_2, 2)], count_from=НАЧАЛО, today=СЕГОДНЯ
    )
    assert [s.code for s in итог.skipped] == [SKIP_UNKNOWN]


def test_книгу_партнёра_openpyxl_разбирает_через_defusedxml() -> None:
    """Файл плана — недоверенный: без defusedxml openpyxl раскрыл бы XML-сущности (бомба)."""
    from openpyxl import xml as openpyxl_xml  # type: ignore[import-untyped]

    assert openpyxl_xml.DEFUSEDXML, "defusedxml не установлен — разбор книги беззащитен"

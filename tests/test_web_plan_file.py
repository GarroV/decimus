"""Файл плана «туда и обратно» на экране (#1168, шаг 3): выгрузка, загрузка, выбор при конфликте.

Двери блока `db` подменены: что кладёт база (замок, ревизия, сторожа) проверяет
`tests/test_db_plan_file.py`. Здесь — заслоны маршрутов (только партнёр, только пока план
ждёт строк, чужой сайт), что дверь получает из файла и от входа, и что видит человек.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import timedelta
from typing import Any

import pytest
from flask.testing import FlaskClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from test_plan_file_core import (
    ЗАПРОС,
    СТРОКА_1,
    СТРОКА_2,
    запрос_плана,
    первая_строка_данных,
    правка,
    сверить,
    строка,
)
from test_web_app import карточка, шапка
from web_harness import ЛОГИН, войти, подменить_двери, собрать

from src.db import action_plans as plans
from src.db import plan_flow as flow
from src.db.plan_file import FileApplied
from src.web import inspections as data
from src.web import plan_sheet

ORIGIN = {"Origin": "http://localhost"}
ФАЙЛ = "88888888-8888-8888-8888-888888888888"
СЕГОДНЯ = plans.today()


def _запрос(*, status: str = "requested", electronic: bool = True, revision: int = 7) -> Any:
    rows = (строка(СТРОКА_1, 1), строка(СТРОКА_2, 2, level="D1", title_ru="Пол чистый"))
    план = запрос_плана(rows, revision)
    return plans.PlanRequest(
        **{**план.__dict__, "status": status, "electronic": electronic, "rows": план.rows}
    )


def _книга(revision: int = 7) -> bytes:
    return plan_sheet.build(_запрос(revision=revision), {}, "ru")


def _заполненная(**ячейки: Any) -> bytes:
    данные = _книга()
    первая = первая_строка_данных(данные)
    return правка(
        данные,
        {
            f"F{первая}": "Переразметим",
            f"G{первая}": "Шеф",
            **{k.replace("_", str(первая)): v for k, v in ячейки.items()},
        },
    )


@pytest.fixture
def зовы(monkeypatch: pytest.MonkeyPatch) -> dict[str, list[dict[str, Any]]]:
    журнал: dict[str, list[dict[str, Any]]] = {"upload": [], "apply": []}
    исход: dict[str, Any] = {"apply": None, "request": _запрос(), "stored": None}

    def загрузить(*args: Any, **kw: Any) -> int:
        журнал["upload"].append({"args": args, **kw})
        return 2

    def применить(*args: Any, **kw: Any) -> FileApplied:
        журнал["apply"].append({"args": args, **kw})
        if исход["apply"] is not None:
            return исход["apply"]  # type: ignore[no-any-return]
        рез = сверить(kw["file_rows"], list(_запрос().rows))
        return FileApplied(рез, len(рез.changes), False, kw["expected_revision"])

    monkeypatch.setattr(plans, "get_request", lambda *_a, **_k: исход["request"])
    monkeypatch.setattr(plans, "upload_version", загрузить)
    monkeypatch.setattr(flow, "apply_file_rows", применить)
    monkeypatch.setattr(plans, "list_requests", lambda **_: plans.PlanList((), False))
    monkeypatch.setattr(plans, "file_for_download", lambda *_a, **_k: исход["stored"])
    monkeypatch.setattr(plans, "fetch_file", lambda ref, **_k: ref)
    монитор = журнал
    монитор["исход"] = [исход]  # type: ignore[assignment]
    return журнал


def _исход(зовы: dict[str, Any]) -> dict[str, Any]:
    return зовы["исход"][0]  # type: ignore[no-any-return]


def _стенд(monkeypatch: pytest.MonkeyPatch, tenant: str) -> Iterator[FlaskClient]:
    monkeypatch.setattr(data, "load_card", lambda *_a, **_k: карточка(шапка(), findings=()))
    подменить_двери(monkeypatch, tenant=tenant, role="admin")
    monkeypatch.setattr("src.web.country.countries", lambda **_: (("GE", 2),))
    with собрать(tenant=tenant).test_client() as client:
        assert войти(client).status_code == 302
        yield client


@pytest.fixture
def партнёр(monkeypatch: pytest.MonkeyPatch, зовы: Any) -> Iterator[FlaskClient]:
    yield from _стенд(monkeypatch, "GE")


@pytest.fixture
def уК(monkeypatch: pytest.MonkeyPatch, зовы: Any) -> Iterator[FlaskClient]:
    yield from _стенд(monkeypatch, "HQ")


def _загрузка(клиент: FlaskClient, байты: bytes, **headers: str) -> Any:
    return клиент.post(
        f"/actions/requests/{ЗАПРОС}/upload",
        data={"file": (io.BytesIO(байты), "plan.xlsx")},
        headers={**ORIGIN, "X-Panel": "1", **headers},
        content_type="multipart/form-data",
    )


# --- выгрузка ----------------------------------------------------------------


def test_партнёр_скачивает_план_из_строк_с_ревизией(партнёр: FlaskClient) -> None:
    ответ = партнёр.get(f"/actions/requests/{ЗАПРОС}/template.xlsx?lang=ru")

    assert ответ.status_code == 200
    разобрано = plan_sheet.parse(ответ.data)
    assert (разобрано.request_id, разобрано.revision) == (ЗАПРОС, 7)
    assert [r.row_id for r in разобрано.rows] == [СТРОКА_1, СТРОКА_2]


def test_уК_план_для_заполнения_не_скачивает(уК: FlaskClient) -> None:
    assert уК.get(f"/actions/requests/{ЗАПРОС}/template.xlsx").status_code == 404


def test_на_согласовании_файл_для_заполнения_не_выдаётся(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    _исход(зовы)["request"] = _запрос(status="on_review")

    assert партнёр.get(f"/actions/requests/{ЗАПРОС}/template.xlsx").status_code == 404


def test_у_файлового_плана_остаётся_прежний_шаблон(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    _исход(зовы)["request"] = _запрос(electronic=False)

    ответ = партнёр.get(f"/actions/requests/{ЗАПРОС}/template.xlsx")

    assert ответ.status_code == 200
    лист = load_workbook(io.BytesIO(ответ.data)).worksheets[0]
    assert лист.cell(row=1, column=10).value is None, "метки электронного шаблона у старого нет"


def test_панель_партнёра_даёт_скачать_и_загрузить_файл_плана(
    партнёр: FlaskClient,
) -> None:
    страница = партнёр.get(f"/actions?tab=plans&id={ЗАПРОС}&lang=ru").get_data(as_text=True)

    assert f"/actions/requests/{ЗАПРОС}/template.xlsx" in страница
    assert "Заполнить файлом" in страница and "Скачать план для заполнения" in страница


# --- загрузка ----------------------------------------------------------------


def test_наш_файл_разбирается_и_строки_идут_в_дверь_с_ревизией_файла(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    ответ = _загрузка(партнёр, _заполненная())

    assert ответ.status_code == 200
    [загрузка] = зовы["upload"]
    assert загрузка["parsed"] is True
    [зов] = зовы["apply"]
    assert зов["args"] == (ЗАПРОС,)
    assert (зов["expected_revision"], зов["tenant"], зов["actor"]) == (7, "GE", ЛОГИН)
    assert [r.row_id for r in зов["file_rows"]] == [СТРОКА_1, СТРОКА_2]
    assert зов["file_rows"][0].measure == "Переразметим"
    текст = ответ.get_data(as_text=True)
    assert "Файл прочитан: строк 2" in текст and "Записано в план строк: 1." in текст


def test_итог_называет_где_нет_меры_и_какой_срок_не_принят(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    позже = СЕГОДНЯ + timedelta(days=400)

    ответ = _загрузка(партнёр, _заполненная(I_=позже))

    текст = ответ.get_data(as_text=True)
    assert "Нет меры или ответственного в строках: 2" in текст
    assert "Строка 1: срок не принят — позже положенного" in текст


def test_строка_без_пары_названа_в_итоге(партнёр: FlaskClient) -> None:
    данные = _заполненная()
    первая = первая_строка_данных(данные)

    ответ = _загрузка(
        партнёр, правка(данные, {f"J{первая + 1}": "99999999-9999-9999-9999-999999999999"})
    )

    assert "Пропущено строк файла без пары: 1" in ответ.get_data(as_text=True)


def test_не_наш_файл_принимается_вложением_и_строки_не_читаются(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    ответ = _загрузка(партнёр, b"%PDF-1.4 plan")

    assert ответ.status_code == 200
    assert зовы["upload"][0]["parsed"] is False and зовы["apply"] == []
    assert "не разобран" in ответ.get_data(as_text=True)


def test_файл_другого_плана_не_читается_как_строки_этого(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    другой = plan_sheet.build(
        plans.PlanRequest(**{**_запрос().__dict__, "id": "22222222-2222-2222-2222-222222222222"}),
        {},
        "ru",
    )

    ответ = _загрузка(партнёр, другой)

    assert зовы["upload"][0]["parsed"] is False and зовы["apply"] == []
    assert "другого плана" in ответ.get_data(as_text=True)


def test_файловому_плану_книгу_не_разбирают(партнёр: FlaskClient, зовы: dict[str, Any]) -> None:
    _исход(зовы)["request"] = _запрос(electronic=False)

    _загрузка(партнёр, _заполненная())

    assert зовы["upload"][0]["parsed"] is False and зовы["apply"] == []


def test_ревизия_разошлась_ничего_не_записано_показана_разница(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    экран = _запрос(revision=9)
    данные = _заполненная()
    разобрано = plan_sheet.parse(данные)
    рез = сверить(разобрано.rows, list(экран.rows))
    _исход(зовы)["apply"] = FileApplied(рез, 0, True, 9)
    _исход(зовы)["request"] = plans.PlanRequest(
        **{
            **экран.__dict__,
            "files": (
                plans.PlanFile(
                    id=ФАЙЛ,
                    version=2,
                    file_name="plan.xlsx",
                    size_bytes=1,
                    content_type="x",
                    uploaded_by="p",
                    uploaded_at=экран.requested_at,
                    verdict=None,
                    comment=None,
                    reviewed_by=None,
                    reviewed_at=None,
                    parsed=True,
                ),
            ),
        }
    )

    ответ = _загрузка(партнёр, данные)

    текст = ответ.get_data(as_text=True)
    assert "План на экране изменили после выгрузки файла" in текст
    assert f'action="/actions/requests/{ЗАПРОС}/file/apply' in текст
    assert f'name="pick_{СТРОКА_1}"' in текст and f'value="{ФАЙЛ}"' in текст
    assert 'name="revision" value="9"' in текст
    assert "Переразметим" in текст


# --- выбор при конфликте -----------------------------------------------------


def _выбор(клиент: FlaskClient, **форма: str) -> Any:
    return клиент.post(
        f"/actions/requests/{ЗАПРОС}/file/apply",
        data={"file": ФАЙЛ, "revision": "9", **форма},
        headers={**ORIGIN, "X-Panel": "1"},
    )


def _файл_в_хранилище(зовы: dict[str, Any], данные: bytes) -> None:
    ref = данные
    _исход(зовы)["stored"] = ref


def test_всё_из_файла_идёт_в_дверь_с_ревизией_показанной_разницы(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    _файл_в_хранилище(зовы, _заполненная())

    ответ = _выбор(партнёр, mode="file")

    assert ответ.status_code == 200
    [зов] = зовы["apply"]
    assert (зов["expected_revision"], зов["take"]) == (9, None)
    assert "Выбор применён: записано строк 1." in ответ.get_data(as_text=True)


def test_выбор_по_строкам_берёт_из_файла_только_отмеченные(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    _файл_в_хранилище(зовы, _заполненная())

    _выбор(партнёр, mode="pick", **{f"pick_{СТРОКА_1}": "file", f"pick_{СТРОКА_2}": "screen"})

    assert зовы["apply"][0]["take"] == frozenset({СТРОКА_1})


def test_всё_как_на_экране_ничего_не_пишет(партнёр: FlaskClient, зовы: dict[str, Any]) -> None:
    ответ = _выбор(партнёр, mode="screen")

    assert зовы["apply"] == []
    assert "Оставлена версия экрана" in ответ.get_data(as_text=True)


def test_пока_выбирали_план_снова_изменили_разница_показана_заново(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    данные = _заполненная()
    _файл_в_хранилище(зовы, данные)
    рез = сверить(plan_sheet.parse(данные).rows, list(_запрос(revision=10).rows))
    _исход(зовы)["apply"] = FileApplied(рез, 0, True, 10)

    ответ = _выбор(партнёр, mode="file")

    текст = ответ.get_data(as_text=True)
    assert "снова изменили" in текст and 'name="revision" value="10"' in текст


def test_файл_другого_плана_в_выборе_не_применяется(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    _файл_в_хранилище(
        зовы,
        plan_sheet.build(
            plans.PlanRequest(
                **{**_запрос().__dict__, "id": "22222222-2222-2222-2222-222222222222"}
            ),
            {},
            "ru",
        ),
    )

    ответ = _выбор(партнёр, mode="file")

    assert зовы["apply"] == []
    assert "Файл не найден или не разбирается" in ответ.get_data(as_text=True)


def test_нет_такого_файла_в_охвате_выбор_отказ(партнёр: FlaskClient, зовы: dict[str, Any]) -> None:
    ответ = _выбор(партнёр, mode="file")

    assert зовы["apply"] == []
    assert "Файл не найден" in ответ.get_data(as_text=True)


def test_выбор_без_ревизии_или_с_мусорным_режимом_отказ(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    _файл_в_хранилище(зовы, _заполненная())

    _выбор(партнёр, mode="file", revision="abc")
    _выбор(партнёр, mode="drop")

    assert зовы["apply"] == []


def test_выбор_уК_отказ(уК: FlaskClient, зовы: dict[str, Any]) -> None:
    ответ = уК.post(
        f"/actions/requests/{ЗАПРОС}/file/apply",
        data={"mode": "file", "file": ФАЙЛ, "revision": "9"},
        headers=ORIGIN,
    )

    assert ответ.status_code == 403 and зовы["apply"] == []


def test_выбор_с_чужого_сайта_отказ(партнёр: FlaskClient, зовы: dict[str, Any]) -> None:
    ответ = партнёр.post(
        f"/actions/requests/{ЗАПРОС}/file/apply",
        data={"mode": "file", "file": ФАЙЛ, "revision": "9"},
        headers={"Origin": "https://evil.example"},
    )

    assert ответ.status_code == 403 and зовы["apply"] == []


def test_у_версии_файла_электронного_плана_метка_разобран_или_нет(
    партнёр: FlaskClient, зовы: dict[str, Any]
) -> None:
    экран = _запрос()

    def версия(n: int, parsed: bool) -> plans.PlanFile:
        return plans.PlanFile(
            id=f"88888888-8888-8888-8888-88888888888{n}", version=n, file_name=f"p{n}.xlsx",
            size_bytes=1, content_type="x", uploaded_by="p", uploaded_at=экран.requested_at,
            verdict=None, comment=None, reviewed_by=None, reviewed_at=None, parsed=parsed,
        )  # fmt: skip

    _исход(зовы)["request"] = plans.PlanRequest(
        **{**экран.__dict__, "files": (версия(1, True), версия(2, False))}
    )

    страница = партнёр.get(f"/actions?tab=plans&id={ЗАПРОС}&lang=ru").get_data(as_text=True)

    assert страница.count("строки прочитаны") == 1 and страница.count("не разобран") == 1

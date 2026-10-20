"""Пункт — меткой «Порядок 13», не кодом `CLN13` (решение владельца 20.10.2026).

Правило одно (`domain.item_label`), поверхности берут его же: бот, веб, письмо, MCP.
Код остаётся связью в данных — здесь проверяется только показ.
"""

from __future__ import annotations

import importlib.util
import io
from datetime import date
from pathlib import Path

import test_web_action_plans
from bot_harness import CHAT_ID
from mcp_checklist_harness import build_edition
from openpyxl import load_workbook  # type: ignore[import-untyped]

from src.bot import view
from src.db.models import FindingRow
from src.domain import Finding, item_label, start_inspection
from src.domain.item_label import ItemTexts
from src.mcp import tools as mcp_tools
from src.report import info_titles
from src.report.letters import Papers
from src.web import plan_template, review

ROOT = Path(__file__).resolve().parent.parent

ТЕКСТЫ = {
    "CLN13": ItemTexts("Порядок", "Tidiness", "Мебель загрязнена", "Furniture is dirty"),
    "MGM08": ItemTexts("Документы", "Documents", "Нет журнала", ""),
    "XYZ05": ItemTexts("Без перевода", "", "Пункт", ""),
}


def test_метка_это_процесс_и_номер_без_ведущих_нулей() -> None:
    assert item_label.label("CLN13", ТЕКСТЫ, "ru") == "Порядок 13"
    assert item_label.label("MGM08", ТЕКСТЫ, "ru") == "Документы 8"


def test_английская_метка_берёт_английское_название_процесса() -> None:
    """Язык экрана — параметр: на английском экране не должно быть русского слова."""
    assert item_label.label("CLN13", ТЕКСТЫ, "en") == "Tidiness 13"
    assert item_label.full("CLN13", ТЕКСТЫ, "en") == "Tidiness 13 — Furniture is dirty"
    assert item_label.full("CLN13", ТЕКСТЫ, "ru") == "Порядок 13 — Мебель загрязнена"


def test_нет_английского_текста_откат_на_русский_а_не_на_код() -> None:
    assert item_label.label("MGM08", ТЕКСТЫ, "en") == "Documents 8"
    assert item_label.full("MGM08", ТЕКСТЫ, "en") == "Documents 8 — Нет журнала"
    assert item_label.label("XYZ05", ТЕКСТЫ, "en") == "Без перевода 5"


def test_служебный_код_называется_словами_без_кода() -> None:
    assert item_label.label("NOTE", ТЕКСТЫ, "en") == "General note"
    assert item_label.label("NOTE", ТЕКСТЫ, "ru") == "Общая заметка"
    assert item_label.full("NOTE", ТЕКСТЫ, "ru") == "Общая заметка"


def test_пункта_нет_в_методике_остаётся_код_и_сбой_виден() -> None:
    assert item_label.label("ZZZ99", ТЕКСТЫ, "ru") == "ZZZ99"
    assert item_label.label("CLN13", None, "ru") == "CLN13"
    assert item_label.full("ZZZ99", ТЕКСТЫ, "ru") == "ZZZ99"


def test_регистр_кода_не_важен() -> None:
    assert item_label.label("cln13", ТЕКСТЫ, "ru") == "Порядок 13"


def test_лист_вычитки_несёт_метку_на_языке_экрана() -> None:
    строки = [
        {
            "id": "CLN13",
            "kind": "violation",
            "process_ru": "Порядок",
            "process_en": "Tidiness",
            "question_ru": "Мебель загрязнена",
            "question_en": "Furniture is dirty",
            "levels": "D1",
            "zones": "*",
        }
    ]
    ру = review.build_sheet(items=строки, zones=[], lang="ru")
    ан = review.build_sheet(items=строки, zones=[], lang="en")

    assert ру.picks[0].label == "Порядок 13"
    assert ан.picks[0].label == "Tidiness 13"
    assert ан.picks[0].code == "CLN13", "код остаётся связью, метка — только показ"


def test_бот_называет_пункт_меткой_на_языке_интерфейса(domain_env: Path) -> None:
    """Список записей перед отчётом: «Порядок 5», «Tidiness 5» — не `CLN05`."""
    start_inspection(CHAT_ID, "Белград 2", "planned", "ru", date="2026-08-21", auditor="Гарро")
    запись = Finding(n=1, code="CLN05", level="D1", zone="hot_kitchen", text="нагар на поду")

    ру = view.record_lines([запись], "ru", chat_id=CHAT_ID)
    ан = view.record_lines([запись], "en", chat_id=CHAT_ID)

    assert "Порядок 5" in ру and "CLN05" not in ру
    assert "Tidiness 5" in ан and "CLN05" not in ан
    assert "Tidiness 5" in view.confirm_line(запись, "en", chat_id=CHAT_ID)
    assert "Порядок 5" in view.changed_line(запись, "ru", chat_id=CHAT_ID)


def test_движок_и_домен_называют_пункт_одинаково() -> None:
    """Движок — отдельный подпроцесс и держит свою копию правила: расхождение — здесь."""
    spec = importlib.util.spec_from_file_location(
        "engine_report_for_labels", ROOT / "engine/report.py"
    )
    assert spec is not None and spec.loader is not None
    движок = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(движок)

    for код, строка in (("CLN13", ТЕКСТЫ["CLN13"]), ("MGM08", ТЕКСТЫ["MGM08"])):
        csv_row = {
            "process_ru": строка.process_ru,
            "process_en": строка.process_en,
        }
        for язык in ("ru", "en"):
            assert движок.item_label(код, csv_row, язык) == item_label.label(код, ТЕКСТЫ, язык)
    assert движок.item_label("ZZZ99", None, "ru") == "ZZZ99"


def test_метка_из_издания_проверки_по_отпечатку(tmp_path: Path) -> None:
    """Тексты берутся из методики ТОЙ версии: издание читается с диска, а не из действующей."""
    живая = tmp_path / "живая"
    версия = build_edition(живая)

    тексты, исход = info_titles.item_texts(версия, papers=Papers(live=живая, store=None))

    assert исход == info_titles.FOUND
    assert item_label.label("CLN01", тексты, "ru") == "Чистота 1"
    assert item_label.full("CLN02", тексты, "en") == "Cleanliness 2 — Walls are clean"


def test_нет_издания_на_машине_метка_остаётся_кодом(tmp_path: Path) -> None:
    тексты, исход = info_titles.item_texts(
        "harness 2026-01-01 deadbeef", papers=Papers(live=None, store=tmp_path / "нет")
    )

    assert исход != info_titles.FOUND
    assert item_label.label("CLN01", тексты, "ru") == "CLN01"


def test_находка_в_mcp_несёт_метку_и_формулировку_а_код_остаётся() -> None:
    запись = FindingRow(
        id="f1", inspection_id="i1", unit_name="Батуми-1", inspection_date=date(2026, 9, 1),
        n=1, code="CLN13", level="D1", zone="hot_kitchen", zone_unusual=False, source="comment",
        lang="ru", text="грязно", comment=None,
    )  # fmt: skip

    ан = mcp_tools._finding(запись, ТЕКСТЫ, "en")
    ру = mcp_tools._finding(запись, ТЕКСТЫ, "ru")
    без = mcp_tools._finding(запись)

    assert ан["code"] == "CLN13" and ру["code"] == "CLN13"
    assert (ан["item_label"], ан["item_title"]) == ("Tidiness 13", "Furniture is dirty")
    assert (ру["item_label"], ру["item_title"]) == ("Порядок 13", "Мебель загрязнена")
    assert без["item_label"] == "CLN13" and без["item_title"] is None


def test_шаблон_плана_показывает_партнёру_метку_а_не_код() -> None:
    запись = FindingRow(
        id="f1", inspection_id="i", unit_name="Batumi-1", inspection_date=date(2026, 9, 1),
        n=1, code="CLN13", level="D2", zone="KITCHEN", zone_unusual=False, source="comment",
        lang="ru", text="Пример записи", comment=None,
    )  # fmt: skip
    книга = plan_template.build(
        test_web_action_plans.запрос(),
        (запись,),
        {"KITCHEN": "Кухня"},
        "ru",
        {"CLN13": "Порядок 13"},
    )

    лист = load_workbook(io.BytesIO(книга)).active
    ячейки = {c for row in лист.iter_rows(values_only=True) for c in row}
    assert "Порядок 13" in ячейки and "CLN13" not in ячейки

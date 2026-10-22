"""Шаблон экшн-плана (#976, ступень 1): строки из нарушений, выдача по охвату."""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

import pytest
from flask.testing import FlaskClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from test_web_action_plans import ЗАПРОС, запрос, зовы  # noqa: F401
from test_web_app import карточка, шапка
from web_harness import войти, подменить_двери, собрать

from src.db import action_plans as plans
from src.db.errors import DbError
from src.db.models import FindingRow
from src.web import inspections as data
from src.web import plan_template
from src.web.action_plans import XLSX_MIME

ДАТА = шапка().inspection_date


def запись(n: int, level: str, code: str, text: str) -> FindingRow:
    return FindingRow(
        id=f"f{n}", inspection_id="i", unit_name="Batumi-1", inspection_date=ДАТА,
        n=n, code=code, level=level, zone="KITCHEN", zone_unusual=False, source="comment",
        lang="ru", text=text, comment=None,
    )  # fmt: skip


ЗАПИСИ = (
    запись(1, "D1", "CLN02", "Пример записи первого класса"),
    запись(2, "R", "ADV01", "Пример рекомендации"),
    запись(3, "D3", "SAF01", "Пример записи третьего класса"),
    запись(4, "D2", "TMP03", "Пример записи второго класса"),
)


def test_в_плане_только_нарушения_и_тяжёлые_сверху() -> None:
    assert [f.code for f in plan_template.plan_rows(ЗАПИСИ)] == ["SAF01", "TMP03", "CLN02"]


def _таблица(книга: bytes) -> list[tuple[Any, ...]]:
    лист = load_workbook(io.BytesIO(книга)).active
    return [tuple(c for c in row) for row in лист.iter_rows(values_only=True)]


def test_книга_с_шапкой_запроса_и_строкой_на_нарушение() -> None:
    строки = _таблица(plan_template.build(запрос(), ЗАПИСИ, {"KITCHEN": "Кухня"}, "ru"))

    assert строки[0][0] == "План исправления нарушений"
    assert ("Прислать план до", None, "11 окт 2026") == строки[5][:3]
    заголовок = next(i for i, r in enumerate(строки) if r[0] == "№")
    assert строки[заголовок][5:] == (
        "Что делаем",
        "Ответственный",
        "Срок",
        "Отметка об исправлении",
    )
    тело = строки[заголовок + 1 :]
    assert [r[:5] for r in тело] == [
        (1, "D3", "Кухня", "SAF01", "Пример записи третьего класса"),
        (2, "D2", "Кухня", "TMP03", "Пример записи второго класса"),
        (3, "D1", "Кухня", "CLN02", "Пример записи первого класса"),
    ]
    assert all(r[5:] == (None, None, None, None) for r in тело)


def test_имя_файла_латиницей() -> None:
    assert plan_template.file_name(запрос()) == "action-plan-Batumi-1-2026-10-03.xlsx"


# --- выдача ------------------------------------------------------------------

АДРЕС = f"/actions/requests/{ЗАПРОС}/template.xlsx"


@pytest.fixture
def партнёр(monkeypatch: pytest.MonkeyPatch, зовы: Any) -> Iterator[FlaskClient]:  # noqa: F811
    monkeypatch.setattr(data, "load_card", lambda *_a, **_k: карточка(шапка(), findings=ЗАПИСИ))
    подменить_двери(monkeypatch, tenant="GE", role="admin")
    monkeypatch.setattr("src.web.country.countries", lambda **_: (("GE", 2),))
    with собрать(tenant="GE").test_client() as client:
        assert войти(client).status_code == 302
        yield client


def test_партнёр_скачивает_шаблон_книгой(партнёр: FlaskClient) -> None:
    ответ = партнёр.get(АДРЕС + "?lang=ru")

    assert ответ.status_code == 200
    assert ответ.mimetype == XLSX_MIME
    assert "attachment" in ответ.headers["Content-Disposition"]
    assert len(_таблица(ответ.data)) > 3


def test_чужой_запрос_404_а_не_книга(партнёр: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(plans, "get_request", lambda *_a, **_k: None)
    assert партнёр.get(АДРЕС).status_code == 404


def test_база_отказала_503_с_причиной(
    партнёр: FlaskClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    def отказ(*_a: Any, **_k: Any) -> None:
        raise DbError("down")

    monkeypatch.setattr(data, "load_card", отказ)
    ответ = партнёр.get(АДРЕС + "?lang=ru")
    assert ответ.status_code == 503
    assert "Шаблон не собрался" in ответ.get_data(as_text=True)


def test_пока_план_ждут_в_панели_сказано_что_ждём_и_есть_шаблон(
    партнёр: FlaskClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(plans, "get_request", lambda *_a, **_k: запрос(status="requested"))
    html = партнёр.get(f"/actions?tab=plans&id={ЗАПРОС}&lang=ru").get_data(as_text=True)
    assert "Что ждём" in html and "до 11 окт 2026" in html
    assert f"/requests/{ЗАПРОС}/template.xlsx" in html


def test_план_на_приёмке_не_просит_загрузить(партнёр: FlaskClient) -> None:
    """Загрузки у плана на приёмке нет — и текста «загрузите ниже» тоже (norma, 19.10)."""
    html = партнёр.get(f"/actions?tab=plans&id={ЗАПРОС}&lang=ru").get_data(as_text=True)
    assert "Что ждём" not in html


def test_строка_с_равно_в_начале_ложится_текстом_а_не_формулой() -> None:
    """Формулировку пишут люди и модель: «=HYPERLINK(…)» не должен стать ссылкой."""
    злая = запись(1, "D1", "X01", '=HYPERLINK("http://example.invalid","x")')
    книга = load_workbook(io.BytesIO(plan_template.build(запрос(), (злая,), {}, "ru"))).active
    ячейка = next(c for row in книга.iter_rows() for c in row if c.value == злая.text)
    assert ячейка.data_type == "s"

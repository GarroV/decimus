"""«Свод пиццерий» девелопмента → справочник `unit_profiles` (D374–D376).

Свод — выгрузка Excel вне git (`workbench/private`). Читаются два листа:
«Детальный свод» (действующие и строящиеся) и «закрытые». Строка пиццерии
опознаётся по стране и имени — кода в своде нет; код Dodo (`dodo_id`)
подбирается по имени из публичного API Dodo (`maps_link.fetch_units`).

Файл — внешние данные: строка без имени или с незнакомой страной
отбрасывается и считается, кривая дата — пустая дата, а не сбой загрузки.
Колонки с именами сотрудников (ответственные, дизайнеры) не переносятся.

Запуск: `python -m src.ratings.svod <путь к xlsx>`.
"""

from __future__ import annotations

import logging
import re
import sys
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from src.db import unit_profiles
from src.db.unit_profiles import Profile

from . import maps_link

logger = logging.getLogger(__name__)

MAIN_SHEET = "Детальный свод"
CLOSED_SHEET = "закрытые"
HEADER_ROW = 1  # строка 0 — итог «Открыто пиццерий», строка 1 — заголовки
MAIN_FIRST_ROW = 3  # строка 2 «Детального свода» — источник данных колонки
CLOSED_FIRST_ROW = 2

COUNTRIES = {
    "Россия": "RU", "Казахстан": "KZ", "Беларусь": "BY", "Турция": "TR",
    "Узбекистан": "UZ", "Нигерия": "NG", "Кыргызстан": "KG", "Румыния": "RO",
    "Объединенные Арабские Эмираты": "AE", "ОАЭ": "AE", "Таджикистан": "TJ",
    "Литва": "LT", "Польша": "PL", "Болгария": "BG", "Монголия": "MN",
    "Эстония": "EE", "Словения": "SI", "Кипр": "CY", "Армения": "AM",
    "Сербия": "RS", "Грузия": "GE", "Черногория": "ME", "Индонезия": "ID",
    "Катар": "QA", "Азербайджан": "AZ", "Испания": "ES", "Ирак": "IQ",
    "Молдова, Республика": "MD", "Мексика": "MX", "Великобритания": "GB",
    "Китай": "CN", "Вьетнам": "VN", "США": "US", "Хорватия": "HR", "Германия": "DE",
}  # fmt: skip

OPEN_STATUSES = {"Ресторан открыт", "Доставка открыта"}
PAUSED_STATUSES = {"Временно закрыта"}

# Колонки, которые лягут в свои поля; всё прочее — в `extra`.
MAIN_COLUMNS = {
    "name": "Пиццерия",
    "partner": "Франчайзи",
    "partner_email": "Почта",
    "address": "Адрес",
    "status": "Статус",
    "revenue_on": "Начало работы (старт выручки в Додо ИС)",
    "delivery_on": "дата поступления выручки по доставке",
    "restaurant_on": "дата поступления выручки по ресторану",
    "closed_on": "Дата Закрытия",
    "city": "Город",
    "country": "Страна",
    "total_area": "Общая площадь",
    "seats": "Количество посадочных мест",
}
CLOSED_COLUMNS = {
    **{k: v for k, v in MAIN_COLUMNS.items() if k not in ("partner", "partner_email")},
    "revenue_on": "Начало работы",
    "delivery_on": "Открытие доставки",
    "restaurant_on": "Открытие ресторана",
}
# Имена сотрудников — не сведения о пиццерии.
DROPPED_COLUMNS = {"Ответственный за проект", "Дизайнер", "Проектировщики"}


def normalize(name: str) -> str:
    """Имя для сравнения: без регистра, пробелов и диакритики («İzmir-1» = «Izmir-1»)."""
    flat = unicodedata.normalize("NFKD", name.casefold())
    return re.sub(r"\s+", "", "".join(ch for ch in flat if not unicodedata.combining(ch)))


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"[​\s]+", " ", str(value)).strip()
    return text or None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _number(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = Decimal(str(value).replace(",", ".").strip())
    except InvalidOperation:
        return None
    return number if number.is_finite() and number >= 0 else None


def _json(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return (value.date() if isinstance(value, datetime) else value).isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return _text(value)


def _stage(status: str, closed_sheet: bool) -> str:
    if closed_sheet:
        return "closed"
    if status in OPEN_STATUSES:
        return "open"
    if status in PAUSED_STATUSES:
        return "paused"
    return "pipeline"


def parse_sheet(
    rows: Sequence[Sequence[Any]], columns: dict[str, str], *, first_row: int, closed: bool
) -> tuple[list[Profile], int]:
    """Строки листа → сведения; второе — сколько строк отброшено."""
    header = [_text(h) for h in rows[HEADER_ROW]] if len(rows) > HEADER_ROW else []
    where = {h: i for i, h in reversed(list(enumerate(header))) if h}
    found: list[Profile] = []
    skipped = 0
    for row in rows[first_row:]:

        def cell(key: str, row: Sequence[Any] = row) -> Any:
            i = where.get(columns[key])
            return row[i] if i is not None and i < len(row) else None

        name, country = _text(cell("name")), COUNTRIES.get(_text(cell("country")) or "")
        status = _text(cell("status"))
        if not name or not country or not status:
            skipped += 1 if any(v not in (None, "") for v in row) else 0
            continue
        seats = _number(cell("seats"))
        mapped = set(columns.values()) | DROPPED_COLUMNS
        extra: dict[str, Any] = {}
        for i, title in enumerate(header):
            if title and title not in mapped and i < len(row) and title not in extra:
                value = _json(row[i])
                if value is not None:
                    extra[title] = value
        found.append(
            Profile(
                country_code=country,
                name=name,
                name_normalized=normalize(name),
                stage=_stage(status, closed),
                status=status,
                city=_text(cell("city")),
                address=_text(cell("address")),
                partner=_text(cell("partner")) if "partner" in columns else None,
                partner_email=_text(cell("partner_email")) if "partner_email" in columns else None,
                restaurant_on=_date(cell("restaurant_on")),
                delivery_on=_date(cell("delivery_on")),
                revenue_on=_date(cell("revenue_on")),
                closed_on=_date(cell("closed_on")),
                total_area=_number(cell("total_area")),
                seats=int(seats) if seats is not None else None,
                extra=extra,
            )
        )
    return found, skipped


def merge(main: Iterable[Profile], closed: Iterable[Profile]) -> list[Profile]:
    """Одна строка на пиццерию: действующая из «Детального свода» главнее
    закрытой (имя переехавшей пиццерии живёт дальше); в пределах листа — первая."""
    taken: dict[tuple[str, str], Profile] = {}
    for profile in [*main, *closed]:
        taken.setdefault((profile.country_code, profile.name_normalized), profile)
    return list(taken.values())


def with_codes(profiles: Sequence[Profile], units: Sequence[maps_link.DodoUnit]) -> list[Profile]:
    """Код Dodo по имени: имя пиццерии в своде и в Dodo IS одно и то же.
    Неоднозначное имя (две пиццерии с одним именем в API) кода не получает."""
    names: dict[str, list[str]] = {}
    for unit in units:
        names.setdefault(normalize(unit.name), []).append(unit.dodo_id)
    out = []
    for p in profiles:
        ids = names.get(normalize(p.name), [])
        out.append(replace(p, dodo_id=ids[0] if len(ids) == 1 else None))
    return out


def unit_codes(
    profiles: Sequence[Profile],
    units: Iterable[tuple[str, str | None, tuple[str, ...], str | None]],
) -> list[tuple[str, str]]:
    """Точке справочника HQ — код Dodo её строки свода: по имени или синониму в
    пределах её страны. Точка без страны или с двумя кандидатами — без кода."""
    codes = {(p.country_code, p.name_normalized): p.dodo_id for p in profiles if p.dodo_id}
    out = []
    for unit_id, country, names, current in units:
        if not country:
            continue
        found = {codes[k] for n in names if (k := (country, normalize(n))) in codes}
        if len(found) == 1 and current not in found:
            out.append((unit_id, found.pop()))
    return out


def read(path: Path) -> list[Profile]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        main, skipped_main = parse_sheet(
            list(book[MAIN_SHEET].iter_rows(values_only=True)),
            MAIN_COLUMNS,
            first_row=MAIN_FIRST_ROW,
            closed=False,
        )
        closed, skipped_closed = parse_sheet(
            list(book[CLOSED_SHEET].iter_rows(values_only=True)),
            CLOSED_COLUMNS,
            first_row=CLOSED_FIRST_ROW,
            closed=True,
        )
    finally:
        book.close()
    logger.info(
        "Свод: действующих %d, закрытых %d, отброшено %d",
        len(main),
        len(closed),
        skipped_main + skipped_closed,
    )
    return merge(main, closed)


def main(argv: Sequence[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    args = list(sys.argv[1:] if argv is None else argv)
    if len(args) != 1:
        print("Использование: python -m src.ratings.svod <путь к xlsx>", file=sys.stderr)
        return 2
    profiles = read(Path(args[0]))
    profiles = with_codes(profiles, maps_link.fetch_units({p.country_code for p in profiles}))
    saved = unit_profiles.save(profiles)
    linked = unit_profiles.set_unit_codes(unit_codes(profiles, unit_profiles.hq_units()))
    logger.info(
        "Свод загружен: пиццерий %d, с кодом Dodo %d, точек справочника связано %d",
        saved,
        sum(1 for p in profiles if p.dodo_id),
        linked,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

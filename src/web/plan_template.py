"""Шаблон экшн-плана — `.xlsx` из записей проверки (#976, ступень 1).

Партнёр получал «загрузите план файлом любого формата» без образца, и руководство
не поняло, чего от него ждут (фидбэк 19.10.2026, пункт 15). Шаблон отвечает на
это сам: строка на каждое нарушение проверки, сверху самые тяжёлые (D3 → D1, как
в карточке), и пустые колонки — что делаем, кто отвечает, к какому сроку.
Формат загрузки по-прежнему любой (D265, D267): шаблон помогает, а не требует.

Сроков по каждому пункту здесь нет намеренно: их считает только движок
(`engine/audit.py score`) в папке проверки, а в базе их нет. Пересчитывать их
здесь значило бы завести второе правило сроков; шаблон ссылается на отчёт PDF,
где они напечатаны движком. Срок подачи самого плана — факт запроса из базы.
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from src.db.action_plans import PlanRequest
from src.db.models import FindingRow

from .dates import human_day
from .texts import t
from .view import by_severity

#: Классы нарушений, по которым ждут план; рекомендации (`R`) и пометки без
#: класса в план не входят — по ним партнёр ничего не обязан.
PLAN_LEVELS = frozenset({"D3", "D2", "D1"})

#: Колонки таблицы: ключ текста и ширина в символах.
_COLUMNS: tuple[tuple[str, int], ...] = (
    ("plans.tpl.col.n", 5),
    ("plans.tpl.col.level", 8),
    ("plans.tpl.col.zone", 18),
    ("plans.tpl.col.code", 10),
    ("plans.tpl.col.finding", 60),
    ("plans.tpl.col.action", 40),
    ("plans.tpl.col.owner", 22),
    ("plans.tpl.col.due", 14),
    ("plans.tpl.col.done", 28),
)
#: Сколько первых колонок заполняет система; остальные — партнёр.
_FILLED = 5

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEAD_FILL = PatternFill("solid", fgColor="EDEDED")
_PARTNER_FILL = PatternFill("solid", fgColor="FFF8E1")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _put(ws: Any, row: int, col: int, value: object) -> Any:
    """Записать значение ячейки; строка — всегда текстом, никогда формулой.

    openpyxl читает строку с «=» в начале как формулу, а формулировку записи,
    название пиццерии и зоны пишут люди и модель: «=HYPERLINK(…)» в шаблоне
    стал бы ссылкой, которую партнёр откроет одним щелчком.
    """
    ячейка = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        ячейка.data_type = "s"
    return ячейка


def plan_rows(findings: Sequence[FindingRow]) -> list[FindingRow]:
    """Нарушения, по которым ждут план, тяжёлые сверху."""
    return [f for f in by_severity(findings) if str(f.level).strip().upper() in PLAN_LEVELS]


def _head(ws: Any, req: PlanRequest, lang: str) -> int:
    """Шапка над таблицей; возвращает номер строки, с которой идёт таблица."""
    итог = "—"
    if req.grade:
        итог = req.grade + (f" · {req.pct:.1f}%" if req.pct is not None else "")
    строки = (
        (t("plans.tpl.unit", lang), req.unit_name),
        (t("plans.tpl.date", lang), human_day(req.inspection_date, lang)),
        (t("plans.tpl.grade", lang), итог),
        (t("plans.tpl.plan_due", lang), human_day(req.due_on, lang)),
    )
    _put(ws, 1, 1, t("plans.tpl.title", lang)).font = _TITLE
    for i, (label, value) in enumerate(строки, start=3):
        _put(ws, i, 1, label).font = _BOLD
        _put(ws, i, 3, value)
    подсказка = 3 + len(строки) + 1
    _put(ws, подсказка, 1, t("plans.tpl.hint", lang))
    return подсказка + 2


def build(
    req: PlanRequest, findings: Sequence[FindingRow], zones: Mapping[str, str], lang: str
) -> bytes:
    """Книга `.xlsx`: шапка запроса и строка на каждое нарушение проверки."""
    wb = Workbook()
    ws = wb.active
    ws.title = t("plans.tpl.sheet", lang)[:31]
    начало = _head(ws, req, lang)
    for col, (key, width) in enumerate(_COLUMNS, start=1):
        ячейка = _put(ws, начало, col, t(key, lang))
        ячейка.font = _BOLD
        ячейка.fill = _HEAD_FILL if col <= _FILLED else _PARTNER_FILL
        ячейка.alignment = _WRAP
        ws.column_dimensions[get_column_letter(col)].width = width
    for i, f in enumerate(plan_rows(findings), start=1):
        строка = начало + i
        значения = (i, f.level, zones.get(f.zone, f.zone), f.code, f.text or "")
        for col, value in enumerate(значения, start=1):
            _put(ws, строка, col, value).alignment = _WRAP
        for col in range(_FILLED + 1, len(_COLUMNS) + 1):
            ws.cell(row=строка, column=col).fill = _PARTNER_FILL
    ws.freeze_panes = ws.cell(row=начало + 1, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def file_name(req: PlanRequest) -> str:
    """Имя файла латиницей и цифрами: так его не испортит ни почта, ни браузер."""
    имя = "".join(c if c.isascii() and c.isalnum() else "-" for c in req.unit_name).strip("-")
    return f"action-plan-{имя or 'unit'}-{req.inspection_date.isoformat()}.xlsx"

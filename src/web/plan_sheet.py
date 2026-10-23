"""Электронный план исправлений — `.xlsx` туда и обратно (шаг 3 эпика #1168).

**Выгрузка.** Книга из текущих строк плана: колонки системы (№, класс, зона, нарушение,
срок по методике) закрыты от правки, жёлтые колонки — партнёру: «что сделаем», «кто
отвечает», «почему», свой срок. В скрытой колонке J — id строки плана, там же в шапке:
метка нашего шаблона, id запроса и ревизия плана на момент выгрузки. По ревизии загрузка
замечает, что план на экране с тех пор поменяли (`plan_flow.apply_file_rows`).

**Разбор — граница системы, файл недоверенный.** Читается как zip-архив библиотекой, формулы
не исполняются (читаются сохранённые значения), внешние ссылки не открываются, а книга с ними
не разбирается вовсе. Размер задаёт вызывающий (предел вложения), здесь — число частей и
распакованный объём архива, число строк и колонок. Сломанный, чужой или слишком большой файл —
`SheetError` с кодом причины: вызывающий принимает его вложением «не разобран».

Сверку строк файла с планом, правила срока и запись делает блок `db` (`plan_file`,
`plan_flow`): здесь только книга и значения из неё.
"""

from __future__ import annotations

import io
import logging
import re
import uuid
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill, Protection  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from src.db.action_plans import PlanRequest
from src.db.plan_file import FileRow

from .dates import human_day
from .texts import t

logger = logging.getLogger(__name__)

#: Метка нашего шаблона: ячейка J1. Версия — в самой метке, новая раскладка — новая метка.
MARKER = "decimus-plan-v1"
#: Подпись столбца id в строке заголовка таблицы: по ней разбор находит, с какой строки данные.
ID_HEADER = "row-id"

#: Раскладка (1-based): колонки партнёра и скрытая служебная.
COL_N, COL_LEVEL, COL_ZONE, COL_TITLE, COL_DUE_SYS = 1, 2, 3, 4, 5
COL_MEASURE, COL_RESPONSIBLE, COL_REASON, COL_DUE_OWN = 6, 7, 8, 9
COL_ID = 10
#: Служебные ячейки в колонке J: метка, id запроса, ревизия.
ROW_MARKER, ROW_REQUEST, ROW_REVISION = 1, 2, 3

#: Пределы разбора. Строк в плане — десятки; пределы ловят мусор и бомбы, а не штатный файл.
MAX_COLS = 30
MAX_DATA_ROWS = 1000
MAX_HEADER_SCAN = 40
MAX_ZIP_MEMBERS = 300
MAX_UNPACKED_BYTES = 64 * 1024 * 1024
MAX_REVISION = 2**31 - 1

#: Коды причин, по которым файл не разобран.
ERR_NOT_XLSX = "not_xlsx"
ERR_TOO_BIG = "too_big"
ERR_EXTERNAL = "external_links"
ERR_NOT_OURS = "not_ours"
ERR_OTHER_PLAN = "other_plan"

_DATE_FORMAT = "yyyy-mm-dd"
_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEAD_FILL = PatternFill("solid", fgColor="EDEDED")
_PARTNER_FILL = PatternFill("solid", fgColor="FFF8E1")
_WRAP = Alignment(wrap_text=True, vertical="top")
_OPEN = Protection(locked=False)
_DOTTED_DATE = re.compile(r"^(\d{1,2})[./](\d{1,2})[./](\d{4})$")
#: Управляющие знаки, которые база не принимает; перенос строки и табуляция остаются.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

#: Колонки таблицы: ключ текста, ширина, колонка партнёра.
_COLUMNS: tuple[tuple[str, int, bool], ...] = (
    ("plans.tpl.col.n", 5, False),
    ("plans.tpl.col.level", 8, False),
    ("plans.tpl.col.zone", 18, False),
    ("plans.tpl.col.finding", 60, False),
    ("plans.sheet.col.due_sys", 16, False),
    ("plans.sheet.col.measure", 44, True),
    ("plans.sheet.col.responsible", 24, True),
    ("plans.sheet.col.reason", 30, True),
    ("plans.sheet.col.due_own", 16, True),
)


class SheetError(Exception):
    """Файл не разобран как шаблон плана. `code` — одна из `ERR_*`."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ParsedSheet:
    """Что прочли из нашего шаблона: чей он, на какой ревизии выгружен, строки."""

    request_id: str
    revision: int
    rows: tuple[FileRow, ...]


# ── Выгрузка ────────────────────────────────────────────────────────────────


def _put(ws: Any, row: int, col: int, value: object) -> Any:
    """Записать значение ячейки; строка — всегда текстом, никогда формулой.

    Мера, ответственный и формулировка нарушения — тексты людей и модели: «=HYPERLINK(…)»
    в книге стал бы ссылкой, которую человек откроет одним щелчком.
    """
    if isinstance(value, str):
        # Из базы текст приходит как положено, но библиотека на знаке, который книга не
        # хранит, падает целиком, и выгрузка не получилась бы из-за одной строки.
        value = _CONTROL.sub("", value)
    ячейка = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        ячейка.data_type = "s"
    return ячейка


def _head(ws: Any, req: PlanRequest, lang: str) -> int:
    """Шапка над таблицей; возвращает номер строки заголовка таблицы."""
    итог = "—"
    if req.grade:
        итог = req.grade + (f" · {req.pct:.1f}%" if req.pct is not None else "")
    строки = (
        (t("plans.tpl.unit", lang), req.unit_name),
        (t("plans.tpl.date", lang), human_day(req.inspection_date, lang)),
        (t("plans.tpl.grade", lang), итог),
        (t("plans.tpl.plan_due", lang), human_day(req.due_on, lang)),
    )
    _put(ws, 1, 1, t("plans.tpl.title", lang)).font = _TITLE
    for i, (label, value) in enumerate(строки, start=3):
        _put(ws, i, 1, label).font = _BOLD
        _put(ws, i, 3, value)
    подсказка = 3 + len(строки) + 1
    _put(ws, подсказка, 1, t("plans.sheet.hint", lang))
    return подсказка + 2


def build(req: PlanRequest, zones: Mapping[str, str], lang: str) -> bytes:
    """Книга `.xlsx` из текущих строк плана: шапка, таблица и служебная колонка J."""
    wb = Workbook()
    ws = wb.active
    ws.title = t("plans.tpl.sheet", lang)[:31]
    заголовок = _head(ws, req, lang)
    _put(ws, ROW_MARKER, COL_ID, MARKER)
    _put(ws, ROW_REQUEST, COL_ID, req.id)
    ws.cell(row=ROW_REVISION, column=COL_ID, value=req.revision)
    _put(ws, заголовок, COL_ID, ID_HEADER)
    for col, (key, width, partner) in enumerate(_COLUMNS, start=1):
        ячейка = _put(ws, заголовок, col, t(key, lang))
        ячейка.font = _BOLD
        ячейка.fill = _PARTNER_FILL if partner else _HEAD_FILL
        ячейка.alignment = _WRAP
        ws.column_dimensions[get_column_letter(col)].width = width
    for строка, row in enumerate(req.rows, start=заголовок + 1):
        значения = {
            COL_N: row.position,
            COL_LEVEL: row.level,
            COL_ZONE: zones.get(row.zone, row.zone),
            COL_TITLE: row.title(lang),
            COL_MEASURE: row.measure or None,
            COL_RESPONSIBLE: row.responsible or None,
            COL_REASON: row.reason or None,
        }
        for col, value in значения.items():
            _put(ws, строка, col, value).alignment = _WRAP
        сроки = (
            (COL_DUE_SYS, row.latest_own_due(req.count_from)),
            (COL_DUE_OWN, row.own_due),
        )
        for col, срок in сроки:
            ячейка = ws.cell(row=строка, column=col, value=срок)
            ячейка.number_format = _DATE_FORMAT
        _put(ws, строка, COL_ID, row.id)
        for col in range(COL_MEASURE, COL_DUE_OWN + 1):
            ячейка = ws.cell(row=строка, column=col)
            ячейка.fill = _PARTNER_FILL
            ячейка.protection = _OPEN
    ws.column_dimensions[get_column_letter(COL_ID)].hidden = True
    ws.freeze_panes = ws.cell(row=заголовок + 1, column=1)
    # Защита без пароля: она не даёт правки случайно, а не прячет файл. Менять ширину
    # колонок и высоту строк можно — тексты длинные.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def file_name(req: PlanRequest) -> str:
    """Имя файла латиницей и цифрами: так его не испортит ни почта, ни браузер."""
    имя = "".join(c if c.isascii() and c.isalnum() else "-" for c in req.unit_name).strip("-")
    return f"action-plan-{имя or 'unit'}-{req.inspection_date.isoformat()}.xlsx"


# ── Разбор ──────────────────────────────────────────────────────────────────


def _check_archive(data: bytes) -> None:
    """Архив разумного размера и без внешних ссылок — до того, как его откроет библиотека."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as архив:
            части = архив.infolist()
    except (zipfile.BadZipFile, OSError, ValueError, NotImplementedError) as exc:
        raise SheetError(ERR_NOT_XLSX) from exc
    if len(части) > MAX_ZIP_MEMBERS or sum(p.file_size for p in части) > MAX_UNPACKED_BYTES:
        raise SheetError(ERR_TOO_BIG)
    if any("externallink" in p.filename.lower() for p in части):
        raise SheetError(ERR_EXTERNAL)


def _clean_text(value: object) -> str:
    """Значение ячейки как текст: без управляющих знаков; число — без хвоста «.0»."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    return _CONTROL.sub("", str(value)).strip()


def _parse_due(value: object) -> tuple[date | None, bool]:
    """Срок из ячейки: `(дата, нечитаема)`. Пусто — дата не названа, это не ошибка."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, False
    if isinstance(value, datetime):
        return value.date(), False
    if isinstance(value, date):
        return value, False
    if isinstance(value, str):
        text = value.strip()
        try:
            if len(text) >= 10 and text[4] == "-":
                return date.fromisoformat(text[:10]), False
            dotted = _DOTTED_DATE.match(text)
            if dotted:
                return date(int(dotted[3]), int(dotted[2]), int(dotted[1])), False
        except ValueError:
            return None, True
    return None, True


def _revision(value: object) -> int:
    if isinstance(value, bool):
        raise SheetError(ERR_NOT_OURS)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, str) and value.strip().isdigit():
        value = int(value.strip())
    if not isinstance(value, int) or not 0 <= value <= MAX_REVISION:
        raise SheetError(ERR_NOT_OURS)
    return value


def _request_id(value: object) -> str:
    try:
        return str(uuid.UUID(_clean_text(value)))
    except ValueError:
        raise SheetError(ERR_NOT_OURS) from None


def _cell(row: tuple[Any, ...], col: int) -> object:
    return row[col - 1] if col <= len(row) else None


def _row_id(value: object) -> str:
    """id строки: нормализованный uuid; мусор остаётся мусором — сверка назовёт его «без пары»."""
    text = _clean_text(value)
    try:
        return str(uuid.UUID(text))
    except ValueError:
        return text[:64]


def _read_rows(ws: Any) -> list[tuple[Any, ...]]:
    """Первые строки листа, не дальше предела: объявленные размеры листа файлу не верим."""
    limit = MAX_HEADER_SCAN + MAX_DATA_ROWS + 1
    return list(ws.iter_rows(min_row=1, max_row=limit, max_col=MAX_COLS, values_only=True))


def parse(data: bytes) -> ParsedSheet:
    """Прочитать наш шаблон плана. Не наш или испорченный — `SheetError`.

    Формулы не исполняются: берётся сохранённое значение ячейки. Читается только первый лист.
    """
    _check_archive(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        logger.info("файл плана не открылся как xlsx: %s", type(exc).__name__)
        raise SheetError(ERR_NOT_XLSX) from exc
    try:
        if not wb.worksheets:
            raise SheetError(ERR_NOT_OURS)
        try:
            rows = _read_rows(wb.worksheets[0])
        except Exception as exc:
            logger.info("лист файла плана не прочитан: %s", type(exc).__name__)
            raise SheetError(ERR_NOT_XLSX) from exc
    finally:
        wb.close()
    return _parsed(rows)


def _parsed(rows: list[tuple[Any, ...]]) -> ParsedSheet:
    def служебная(номер: int) -> object:
        return _cell(rows[номер - 1], COL_ID) if len(rows) >= номер else None

    if _clean_text(служебная(ROW_MARKER)) != MARKER:
        raise SheetError(ERR_NOT_OURS)
    request_id = _request_id(служебная(ROW_REQUEST))
    revision = _revision(служебная(ROW_REVISION))
    заголовок = next(
        (
            i
            for i, row in enumerate(rows[:MAX_HEADER_SCAN], start=1)
            if _clean_text(_cell(row, COL_ID)) == ID_HEADER
        ),
        None,
    )
    if заголовок is None:
        raise SheetError(ERR_NOT_OURS)
    return ParsedSheet(request_id, revision, _data_rows(rows, заголовок))


def _data_rows(rows: list[tuple[Any, ...]], header: int) -> tuple[FileRow, ...]:
    found: list[FileRow] = []
    for line, row in enumerate(rows[header:], start=header + 1):
        row_id = _row_id(_cell(row, COL_ID))
        measure = _clean_text(_cell(row, COL_MEASURE))
        responsible = _clean_text(_cell(row, COL_RESPONSIBLE))
        reason = _clean_text(_cell(row, COL_REASON))
        due, unreadable = _parse_due(_cell(row, COL_DUE_OWN))
        if not (row_id or measure or responsible or reason or due or unreadable):
            continue
        if len(found) >= MAX_DATA_ROWS:
            raise SheetError(ERR_TOO_BIG)
        found.append(FileRow(line, row_id, measure, responsible, reason, due, unreadable))
    return tuple(found)
